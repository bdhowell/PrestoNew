import asyncio
import io
import os
import tempfile
import unittest
import uuid
import zipfile
from dataclasses import replace
from pathlib import Path
from unittest.mock import AsyncMock, Mock, patch
from urllib.parse import quote

import psycopg
from fastapi import Depends, FastAPI
from fastapi.staticfiles import StaticFiles
from fastapi.testclient import TestClient
from google.api_core.exceptions import ServiceUnavailable
from starlette.requests import HTTPConnection

import content
import source_processing
import posters
import test_posters
from routers import content as routes
from test_posters import DRAFT, ROOT

ITEM = content.Source(1, "<paper>.pdf", "application/pdf", 12, "2026-10-25")


class ContentRouteTests(unittest.TestCase):
    def setUp(self) -> None:
        self.uid = 7
        self.session = {"poster_id": 11}
        self.posters = AsyncMock()
        self.posters.get.return_value = DRAFT
        self.sources = AsyncMock()
        self.sources.list_sources.return_value = []
        self.sources.add.return_value = ITEM
        self.sources.completion.return_value = (0, 0)
        for target, value in (("posters.store", self.posters), ("content.store", self.sources)):
            mocked = patch(target, return_value=value)
            mocked.start()
            self.addCleanup(mocked.stop)

        async def identity(connection: HTTPConnection) -> None:
            connection.state.user_id = self.uid
            connection.state.user = {"user-name": "<Ada>"}
            connection.scope["session"] = self.session

        application = FastAPI(dependencies=[Depends(identity)])
        application.include_router(routes.router)
        application.mount("/static", StaticFiles(directory=ROOT / "static"))
        self.client = TestClient(application)

    def upload(self, data=b"%PDF-1.7\n", filename="paper.pdf", **kwargs):
        return self.client.post("/content/upload", content=data, headers={
            "Content-Type": "application/octet-stream", "X-Filename": quote(filename),
            **kwargs,
        })

    def test_empty_page_shows_introduction_and_all_supported_extensions(self) -> None:
        response = self.client.get("/content/upload")
        self.assertEqual(response.status_code, 200)
        self.assertIn("start by uploading your source content", response.text)
        self.assertIn("Apple Numbers, Pages, and Keynote", response.text)
        self.assertIn("Markdown (MD)", response.text)
        self.assertIn('id="items-list"', response.text)
        self.assertIn("/static/javascript/pages/content-upload.js", response.text)
        for suffix in content.FORMATS:
            self.assertIn(suffix, response.text)
        self.sources.list_sources.assert_awaited_once_with(7, 11)

    def test_existing_files_escape_names_and_do_not_repeat_first_upload_intro(self) -> None:
        self.sources.list_sources.return_value = [ITEM]
        response = self.client.get("/content/upload")
        self.assertIn("&lt;paper&gt;.pdf", response.text)
        self.assertNotIn("start by uploading your source content", response.text)
        self.assertNotIn("Rich Text Format", response.text)

    def test_reduced_conversion_notice_is_visible_on_page_and_upload_response(self) -> None:
        with patch.dict(os.environ, {"DOCLING_ENABLE_ENRICHMENT": "false"}):
            page = self.client.get("/content/upload")
            response = self.upload()
        self.assertIn(source_processing.REDUCED_FEATURE_NOTICE, page.text)
        self.assertEqual(response.status_code, 201)
        self.assertIn(source_processing.REDUCED_FEATURE_NOTICE, response.json()["say"])
        with patch.dict(os.environ, {"DOCLING_ENABLE_ENRICHMENT": "true"}):
            self.assertNotIn(source_processing.REDUCED_FEATURE_NOTICE, self.client.get("/content/upload").text)

    def test_upload_returns_metadata_and_passes_owner_checked_poster(self) -> None:
        response = self.upload(filename="résumé.pdf")
        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.json()["item"]["sid"], 1)
        self.assertNotIn("storage_uri", response.text)
        self.sources.add.assert_awaited_once_with(7, 11, "résumé.pdf", b"%PDF-1.7\n")
        self.posters.get.assert_awaited_once_with(7, 11)

    def test_upload_requires_identity_owned_selected_draft(self) -> None:
        self.uid = None
        response = self.client.get("/content/upload", follow_redirects=False)
        self.assertEqual(response.headers["location"], "/account/signin")
        self.assertEqual(self.upload().status_code, 401)
        self.uid = 7
        self.session.clear()
        self.assertEqual(self.upload().status_code, 409)
        self.session["poster_id"] = 11
        self.posters.get.return_value = None
        self.assertEqual(self.upload().status_code, 404)
        self.posters.get.return_value = replace(DRAFT, purchased=True)
        self.assertEqual(self.upload().status_code, 409)
        self.sources.add.assert_not_awaited()

    def test_invalid_requests_are_rejected_before_storage(self) -> None:
        self.assertEqual(self.client.post("/content/upload", json={}).status_code, 415)
        self.assertEqual(self.upload(Origin="https://evil.example").status_code, 403)
        for name in ("", "../paper.pdf", "folder\\paper.pdf", "file.exe", "a\n.pdf", "a" * 260 + ".pdf", "file.rtf"):
            self.assertEqual(self.upload(filename=name).status_code, 422)
        self.sources.add.assert_not_awaited()

    def test_streamed_upload_limit_is_enforced_without_content_length(self) -> None:
        with patch.dict(os.environ, {"POSTER_UPLOAD_MAX_MB": "1"}):
            self.assertEqual(self.upload(data=(b"x" * 524288 for _ in range(3))).status_code, 413)
        self.sources.add.assert_not_awaited()

    def test_validation_and_storage_failures_are_explicit(self) -> None:
        self.sources.add.side_effect = ValueError("Empty files cannot be uploaded.")
        self.assertEqual(self.upload(data=b"").status_code, 422)
        for error in (psycopg.OperationalError("down"), OSError("disk full"), ServiceUnavailable("down")):
            self.sources.add.side_effect = error
            with self.assertLogs("routers.content", "ERROR"):
                self.assertEqual(self.upload().status_code, 503)
        self.sources.list_sources.side_effect = psycopg.OperationalError("down")
        with self.assertLogs("routers.content", "ERROR"):
            self.assertEqual(self.client.get("/content/upload").status_code, 503)

    def test_finish_checks_empty_no_text_and_exact_configured_threshold(self) -> None:
        with patch.dict(os.environ, {"SOURCE_TEXT_MIN_LENGTH": "4000"}):
            for count, length, message in (
                (0, 0, "you can&rsquo;t have a poster without any content!"),
                (1, 0, "I can&rsquo;t magic a complete poster out of thin air."),
                (1, 3999, "I&rsquo;m unable to conjure enough content to create a useful poster."),
            ):
                self.sources.completion.return_value = (count, length)
                response = self.client.post("/content/finished", json={})
                self.assertEqual(response.status_code, 409)
                self.assertNotIn("redirect", response.json())
                self.assertIn("&lt;Ada&gt;", response.json()["say"])
                self.assertIn(message, response.json()["say"])
            for length in (4000, 4001):
                self.sources.completion.return_value = (1, length)
                response = self.client.post("/content/finished", json={})
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.json(), {"ok": True, "redirect": "/content/elucidate"})
        self.sources.completion.assert_awaited_with(7, 11)

    def test_finish_requires_same_origin_json_and_editable_owner(self) -> None:
        self.assertEqual(self.client.post("/content/finished", json={}, headers={"Origin": "https://evil.example"}).status_code, 403)
        self.assertEqual(self.client.post("/content/finished", content="").status_code, 415)
        self.posters.get.return_value = replace(DRAFT, purchased=True)
        self.assertEqual(self.client.post("/content/finished", json={}).status_code, 409)
        self.sources.completion.assert_not_awaited()

    def test_finished_default_minimum_is_four_thousand_and_errors_are_visible(self) -> None:
        with patch.dict(os.environ, {}, clear=True):
            self.sources.completion.return_value = (1, 3999)
            self.assertEqual(self.client.post("/content/finished", json={}).status_code, 409)
            self.sources.completion.return_value = (1, 4000)
            self.assertEqual(self.client.post("/content/finished", json={}).status_code, 200)
        self.sources.completion.side_effect = psycopg.OperationalError("down")
        with self.assertLogs("routers.content", "ERROR"):
            self.assertEqual(self.client.post("/content/finished", json={}).status_code, 503)

    def test_processing_service_configuration_errors_are_reported_explicitly(self) -> None:
        self.sources.add.side_effect = source_processing.ProcessingError("DOCLING_SERVICE_URL is required")
        with self.assertLogs("routers.content", "ERROR"):
            response = self.upload()
        self.assertEqual(response.status_code, 503)
        self.assertIn("DOCLING_SERVICE_URL", response.json()["detail"])

    def test_size_warning_counts_exact_threshold_and_first_description_line(self) -> None:
        small = replace(ITEM, width=299, height=1000, description="<Logo>\nsecond line")
        large = replace(ITEM, sid=2, width=300, height=300)
        self.assertTrue(small.ill_sized)
        self.assertFalse(large.ill_sized)
        self.assertFalse(ITEM.ill_sized)
        self.assertIn("One file will likely", routes.size_warning([small, large]))
        self.assertIn("Two files will likely", routes.size_warning([small, small]))
        self.sources.list_sources.return_value = [small, large]
        response = self.client.get("/content/upload")
        self.assertIn('class="ill-sized"', response.text)
        self.assertIn("&lt;Logo&gt;", response.text)
        self.assertNotIn("second line", response.text)
        self.assertIn("/content/1/thumbnail", response.text)
        self.assertIn('id="finish-upload"', response.text)

    def test_thumbnail_is_owner_checked_and_reencoded_as_private_png(self) -> None:
        from PIL import Image
        buffer = io.BytesIO()
        Image.new("RGB", (1000, 500)).save(buffer, format="JPEG")
        self.sources.graphic.return_value = buffer.getvalue()
        response = self.client.get("/content/1/thumbnail")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers["cache-control"], "private, no-store")
        with Image.open(io.BytesIO(response.content)) as image:
            self.assertEqual(image.size, (160, 80))
            self.assertEqual(image.format, "PNG")
        self.sources.graphic.assert_awaited_once_with(7, 11, 1)
        self.sources.graphic.return_value = None
        self.assertEqual(self.client.get("/content/99/thumbnail").status_code, 404)


class SourceFileTests(unittest.TestCase):
    def test_hash_handles_text_and_bytes_consistently(self) -> None:
        self.assertEqual(content.source_hash("résumé"), content.source_hash("résumé".encode()))
        self.assertEqual(len(content.source_hash(b"research")), 64)

    def test_all_supported_formats_validate_and_reject_mismatches(self) -> None:
        headers = {
            ".pdf": b"%PDF-1.7", ".png": b"\x89PNG\r\n\x1a\n",
            ".jpg": b"\xff\xd8\xff", ".jpeg": b"\xff\xd8\xff",
            ".tif": b"II*\0", ".tiff": b"MM\0*",
            ".md": b"# Research",
        }
        for suffix in content.FORMATS:
            if suffix in headers:
                data = headers[suffix]
            else:
                buffer = io.BytesIO()
                with zipfile.ZipFile(buffer, "w") as archive:
                    prefix = {".docx": "word/", ".pptx": "ppt/", ".xlsx": "xl/", ".xlsm": "xl/"}.get(suffix, "Index/")
                    archive.writestr(prefix + "document.xml", "document")
                data = buffer.getvalue()
            with self.subTest(suffix=suffix):
                self.assertEqual(content.validate_file("source" + suffix, data), content.FORMATS[suffix])
                with self.assertRaises(ValueError):
                    content.validate_file("source" + suffix, b"")
                if suffix != ".md":
                    with self.assertRaises(ValueError):
                        content.validate_file("source" + suffix, b"not this file format")
        with self.assertRaises(ValueError):
            content.validate_file("file.md", b"\xff")
        with self.assertRaises(ValueError):
            content.validate_file("file.md", b"binary\0text")

    def test_local_files_are_uuid_named_and_sharded_across_ten_directories(self) -> None:
        with tempfile.TemporaryDirectory() as directory, patch.dict(os.environ, {"SOURCES_DIR": directory, "GCS_BUCKET": ""}):
            files = content.SourceFiles()
            for number in range(10):
                fileid = uuid.UUID(int=number)
                uri = files.put(fileid, b"source", "application/pdf")
                self.assertEqual(Path(uri).relative_to(Path(directory).resolve()).as_posix(), f"{number}/{fileid}")
                self.assertEqual(Path(uri).read_bytes(), b"source")
                files.delete(uri)
                self.assertFalse(Path(uri).exists())

    def test_gcs_uses_configured_bucket_and_uuid_key(self) -> None:
        client = Mock()
        with patch.dict(os.environ, {"GCS_BUCKET": "source-test"}), patch.object(content.storage, "Client", return_value=client):
            files = content.SourceFiles()
            fileid = uuid.UUID(int=17)
            uri = files.put(fileid, b"source", "application/pdf")
            self.assertEqual(uri, f"gs://source-test/7/{fileid}")
            client.bucket.assert_called_with("source-test")
            client.bucket().blob.assert_called_with(f"7/{fileid}")
            client.bucket().blob().upload_from_string.assert_called_once_with(
                b"source", content_type="application/pdf", if_generation_match=0, timeout=60,
            )
            files.delete(uri)
            client.bucket().blob().delete.assert_called_once_with(timeout=60)


class SourceStoreTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self) -> None:
        for name, value in (
            ("extract", Mock(return_value=content.processing.Processed(text="text"))),
            ("summarize", AsyncMock(return_value="Summary")),
        ):
            mocked = patch.object(content.processing, name, value)
            mocked.start()
            self.addCleanup(mocked.stop)

    async def test_list_sources_scopes_owner_and_poster_and_orders_chronologically(self) -> None:
        store = content.SourceStore()
        connection = AsyncMock()
        connection.execute.return_value.fetchall.return_value = [
            (1, "file.md", "text/markdown", 4, "now"),
        ]
        with patch.object(store, "_connect", AsyncMock(return_value=connection)):
            items = await store.list_sources(7, 11)
        self.assertEqual(items, [content.Source(1, "file.md", "text/markdown", 4, "now")])
        statement, params = connection.execute.call_args.args
        self.assertIn("p.uid = %s AND p.pid = %s", statement)
        self.assertIn("ORDER BY s.added, s.sid", statement)
        self.assertEqual(params, (7, 11))

    async def test_failed_database_insert_removes_saved_file(self) -> None:
        store = content.SourceStore()
        connection = AsyncMock()
        owner, poster = AsyncMock(), AsyncMock()
        owner.fetchone.return_value = (7,)
        poster.fetchone.return_value = (11,)
        connection.execute.side_effect = [owner, poster, psycopg.OperationalError("down")]
        files = Mock()
        files.put.return_value = "/sources/1/test"
        with patch.object(store, "_connect", AsyncMock(return_value=connection)), patch.object(content, "SourceFiles", return_value=files):
            with self.assertRaises(psycopg.OperationalError):
                await store.add(7, 11, "file.md", b"text")
        files.delete.assert_called_once_with("/sources/1/test")

    async def test_successful_insert_preserves_file_and_records_source_metadata(self) -> None:
        store = content.SourceStore()
        connection = AsyncMock()
        owner, poster, inserted = AsyncMock(), AsyncMock(), AsyncMock()
        owner.fetchone.return_value = (7,)
        poster.fetchone.return_value = (11,)
        inserted.fetchone.return_value = (1, "file.md", "text/markdown", 4, "now")
        connection.execute.side_effect = [owner, poster, inserted]
        files = Mock()
        files.put.return_value = "/sources/1/test"
        with patch.object(store, "_connect", AsyncMock(return_value=connection)), patch.object(content, "SourceFiles", return_value=files):
            result = await store.add(7, 11, "file.md", b"text")
        self.assertEqual(result.filename, "file.md")
        files.delete.assert_not_called()
        statement, params = connection.execute.await_args_list[2].args
        self.assertEqual(params[:6], (11, "text", None, "file.md", "text/markdown", 4))
        self.assertIsInstance(params[6], uuid.UUID)
        self.assertEqual(params[7:11], (content.source_hash(b"text"), "/sources/1/test", "text", "Summary"))

    async def test_foreign_or_purchased_poster_never_writes_a_file(self) -> None:
        store = content.SourceStore()
        connection = AsyncMock()
        owner, poster = AsyncMock(), AsyncMock()
        owner.fetchone.return_value = (7,)
        poster.fetchone.return_value = None
        connection.execute.side_effect = [owner, poster]
        files = Mock()
        with patch.object(store, "_connect", AsyncMock(return_value=connection)), patch.object(content, "SourceFiles", return_value=files):
            with self.assertRaises(ValueError):
                await store.add(7, 99, "file.md", b"text")
        files.put.assert_not_called()

    async def test_identical_graphics_reuse_existing_poster_scoped_record(self) -> None:
        store = content.SourceStore()
        connection = AsyncMock()
        connection.execute.return_value.fetchone.return_value = (
            1, "figure.png", "image/png", 10, "now", None, "figure", 300, 400, "image",
        )
        files = Mock()
        asset = source_processing.Asset("figure.png", b"image", "image/png", "image", width=300, height=400)
        result = await store._insert_source(connection, files, [], 11, asset, 7)
        self.assertEqual(result.sid, 1)
        statement, params = connection.execute.call_args.args
        self.assertIn("s.pid = %s AND s.hash = %s", statement)
        self.assertEqual(params, (11, content.source_hash(b"image")))
        files.put.assert_not_called()

    async def test_logo_replacement_uses_pixel_area_and_keeps_smaller_as_image(self) -> None:
        for old_area, expected_kind in ((100, "logo"), (1000000, "image"), (90000, "image")):
            with self.subTest(old_area=old_area):
                store = content.SourceStore()
                connection = AsyncMock()
                missing, old, inserted, updated = AsyncMock(), AsyncMock(), AsyncMock(), AsyncMock()
                missing.fetchone.return_value = None
                old.fetchone.return_value = (9, old_area)
                inserted.fetchone.return_value = (
                    10, "logo.png", "image/png", 5, "now", None, "A logo", 300, 300, expected_kind,
                )
                connection.execute.side_effect = (
                    [missing, old, updated, inserted] if expected_kind == "logo"
                    else [missing, old, inserted]
                )
                files = Mock()
                files.put.return_value = "/sources/logo"
                asset = source_processing.Asset("logo.png", b"image", "image/png", "image",
                                                description="A logo", width=300, height=300)
                result = await store._insert_source(connection, files, [], 11, asset, None)
                self.assertEqual(result.kind, expected_kind)
                self.assertEqual(connection.execute.await_args_list[-1].args[1][1], expected_kind)
                if expected_kind == "logo":
                    self.assertIn("UPDATE source SET kind = 'image'", connection.execute.await_args_list[2].args[0])

    async def test_children_record_document_parent_and_all_files_rollback_on_failure(self) -> None:
        store = content.SourceStore()
        connection = AsyncMock()
        owner, poster, inserted, missing = AsyncMock(), AsyncMock(), AsyncMock(), AsyncMock()
        owner.fetchone.return_value = (7,)
        poster.fetchone.return_value = (11,)
        inserted.fetchone.return_value = (1, "file.md", "text/markdown", 4, "now")
        missing.fetchone.return_value = None
        connection.execute.side_effect = [owner, poster, inserted, missing, psycopg.OperationalError("down")]
        files = Mock()
        files.put.side_effect = ["/sources/document", "/sources/child"]
        processed = source_processing.Processed(text="text", assets=[
            source_processing.Asset("formula.txt", b"e=mc^2", "text/plain", "image", text="e=mc^2"),
        ])
        with patch.object(content.processing, "extract", return_value=processed), \
                patch.object(store, "_connect", AsyncMock(return_value=connection)), \
                patch.object(content, "SourceFiles", return_value=files):
            with self.assertRaises(psycopg.OperationalError):
                await store.add(7, 11, "file.md", b"text")
        child_params = connection.execute.await_args_list[-1].args[1]
        self.assertEqual(child_params[2], 1)
        self.assertEqual([call.args[0] for call in files.delete.call_args_list], ["/sources/document", "/sources/child"])

    async def test_spreadsheet_markup_does_not_count_toward_text_threshold(self) -> None:
        connection = AsyncMock()
        connection.execute.return_value.fetchall.return_value = [
            ("<table><tr><td>A &amp; B</td></tr></table>", content.FORMATS[".xlsx"][0]),
            ("research", "text/markdown"), (None, "image/png"),
        ]
        store = content.SourceStore()
        with patch.object(store, "_connect", AsyncMock(return_value=connection)):
            self.assertEqual(await store.completion(7, 11), (3, 13))

    async def test_hierarchy_places_children_under_parents_in_chronological_order(self) -> None:
        connection = AsyncMock()
        connection.execute.return_value.fetchall.return_value = [
            (1, "first.pdf", "application/pdf", 1, "1", None, None, None, None, "mixed"),
            (2, "second.pdf", "application/pdf", 1, "2", None, None, None, None, "mixed"),
            (3, "child.png", "image/png", 1, "3", 1, None, 300, 300, "image"),
        ]
        store = content.SourceStore()
        with patch.object(store, "_connect", AsyncMock(return_value=connection)):
            items = await store.list_sources(7, 11)
        self.assertEqual([(item.sid, item.depth) for item in items], [(1, 0), (3, 1), (2, 0)])


class ProcessingTests(unittest.IsolatedAsyncioTestCase):
    async def test_spreadsheet_cells_and_sheet_names_are_html_escaped(self) -> None:
        from openpyxl import Workbook
        workbook = Workbook()
        workbook.active.title = "<Research>"
        workbook.active.append(["<script>bad</script>", "Data & results", 42])
        buffer = io.BytesIO()
        workbook.save(buffer)
        for extension in (".xlsx", ".xlsm"):
            result = source_processing.extract("sheet" + extension, buffer.getvalue())
            self.assertIn("&lt;Research&gt;", result.text)
            self.assertIn("&lt;script&gt;bad&lt;/script&gt;", result.text)
            self.assertIn("Data &amp; results", result.text)
            self.assertNotIn("<script>", result.text)

    async def test_docling_extracts_text_assets_and_enables_descriptions_and_math(self) -> None:
        from PIL import Image
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "figure.png"
            Image.new("RGB", (299, 600)).save(path)
            manifest = {
                "output_directory": directory, "warnings": [],
                "source_texts": [{"id": "caption", "text": "Figure 1"}],
                "elements": [{
                    "kind": "picture", "caption_refs": ["caption"], "annotations": [{"text": "A graph"}],
                    "text": None, "assets": ["figure.png"],
                }, {
                    "kind": "formula", "caption_refs": [], "annotations": [],
                    "text": r"E=mc^2", "assets": [],
                }],
            }
            with patch.dict(os.environ, {"DOCLING_SERVICE_URL": "https://converter.invalid"}), \
                    patch.object(source_processing.docxtractor, "extract_document", return_value=manifest) as convert:
                result = source_processing.extract("research.pdf", b"%PDF-1.7")
            self.assertEqual(result.text, "Figure 1\n\nE=mc^2")
            self.assertEqual(result.assets[0].width, 299)
            self.assertEqual(result.assets[0].description, "Figure 1\nA graph")
            self.assertEqual(result.assets[1].text, "E=mc^2")
            self.assertTrue(convert.call_args.kwargs["describe_pictures"])
            self.assertTrue(convert.call_args.kwargs["extract_chart_data"])
            self.assertTrue(convert.call_args.kwargs["enrich_formulas"])

    async def test_apple_documents_always_use_docling_native_conversion(self) -> None:
        manifest = {"output_directory": "/tmp", "warnings": [], "source_texts": [], "elements": []}
        with patch.dict(os.environ, {"DOCLING_SERVICE_URL": "https://converter.invalid"}), \
                patch.object(source_processing.docxtractor, "extract_document", return_value=manifest) as convert:
            for suffix in (".pages", ".key"):
                source_processing.extract("research" + suffix, b"archive")
                self.assertTrue(convert.call_args.kwargs["native_iwork"])
                self.assertEqual(convert.call_args.kwargs["filename"], "research" + suffix)

    def test_reduced_conversion_is_explicit_and_retains_source_text(self) -> None:
        manifest = {
            "output_directory": "/tmp", "warnings": [],
            "source_texts": [{"id": "paragraph", "text": "Research evidence"}], "elements": [],
        }
        with patch.dict(os.environ, {
            "DOCLING_SERVICE_URL": "https://converter.invalid", "DOCLING_ENABLE_ENRICHMENT": "false",
        }), patch.object(source_processing.docxtractor, "extract_document", return_value=manifest) as convert:
            result = source_processing.extract("research.pdf", b"%PDF-1.7")
        for option in ("describe_pictures", "extract_chart_data", "enrich_formulas"):
            self.assertIs(convert.call_args.kwargs[option], False)
        self.assertEqual(result.text, "Research evidence")
        self.assertIn(source_processing.REDUCED_FEATURE_NOTICE, result.warnings)

    def test_enrichment_defaults_to_enabled_and_rejects_invalid_settings(self) -> None:
        with patch.dict(os.environ, {}, clear=True):
            self.assertTrue(source_processing.enrichment_enabled())
        with patch.dict(os.environ, {"DOCLING_ENABLE_ENRICHMENT": "off"}):
            with self.assertRaisesRegex(source_processing.ProcessingError, "must be true or false"):
                source_processing.enrichment_enabled()

    async def test_rtf_is_rejected_by_validation_and_processing(self) -> None:
        self.assertNotIn(".rtf", content.FORMATS)
        with self.assertRaises(ValueError):
            content.validate_file("file.rtf", b"{\\rtf1 text}")
        with self.assertRaises(ValueError):
            source_processing.extract("file.rtf", b"{\\rtf1 text}")

    async def test_numbers_requires_actual_table_data_and_produces_html(self) -> None:
        manifest = {"output_directory": "/tmp", "warnings": [], "source_texts": [], "elements": []}
        with patch.dict(os.environ, {"DOCLING_SERVICE_URL": "https://converter.invalid"}), \
                patch.object(source_processing.docxtractor, "extract_document", return_value=manifest):
            with self.assertRaises(source_processing.ProcessingError):
                source_processing.extract("sheet.numbers", b"archive")
            manifest["elements"] = [{
                "kind": "table", "caption_refs": [], "annotations": [], "assets": [], "text": "",
                "data": {"table_cells": [
                    {"start_row_offset_idx": 0, "start_col_offset_idx": 0, "text": "<Research>"},
                ]},
            }]
            result = source_processing.extract("sheet.numbers", b"archive")
            self.assertEqual(result.text, "<table><tr><td>&lt;Research&gt;</td></tr></table>")

    def test_refused_docling_requests_explain_instance_and_key_checks(self) -> None:
        from docling.service_client.exceptions import ServiceError
        refused = ServiceError("Task submission failed.", status_code=403)
        with patch.dict(os.environ, {"DOCLING_SERVICE_URL": "https://docling.example"}), \
                patch.object(source_processing.docxtractor, "extract_document", side_effect=refused):
            with self.assertRaises(source_processing.ProcessingError) as raised:
                source_processing.extract("file.md", b"# Research")
        self.assertIn("refused the request (HTTP 403)", str(raised.exception))
        self.assertIn("instance is enabled", str(raised.exception))
        self.assertIs(raised.exception.__cause__, refused)

    async def test_missing_services_fail_explicitly(self) -> None:
        with patch.dict(os.environ, {"DOCLING_SERVICE_URL": "", "OPENAI_API_KEY": ""}):
            with self.assertRaises(source_processing.ProcessingError):
                source_processing.extract("file.md", b"# Research")
            with self.assertRaises(source_processing.ProcessingError):
                await source_processing.summarize("research")

    async def test_upload_chat_finished_replies_invoke_the_same_validated_finish_command(self) -> None:
        from routers.wikkit import finished_uploading, upload_workflow
        for reply in ("Finished", "I'm finished", "I’m finished", "Done", "I'm done", "All done!"):
            with self.subTest(reply=reply):
                self.assertTrue(finished_uploading(reply))
        for reply in ("No", "yes", "What formats can you read?"):
            with self.subTest(reply=reply):
                self.assertFalse(finished_uploading(reply))
        replies = asyncio.Queue()
        for reply in ("I'm all done", "What formats can you read?"):
            replies.put_nowait(reply)
        async def next_reply(**_: object) -> str:
            return await replies.get()

        bot = AsyncMock()
        bot.request_response.side_effect = next_reply
        task = asyncio.create_task(upload_workflow(bot))
        try:
            for _ in range(20):
                if bot.handle_message.await_count:
                    break
                await asyncio.sleep(0)
            bot.request_response.assert_awaited_with(remind=False)
            bot.send.assert_awaited_once_with({"command": {"does": "finished-uploading"}})
            bot.handle_message.assert_awaited_once_with("What formats can you read?")
        finally:
            task.cancel()
            with self.assertRaises(asyncio.CancelledError):
                await task

    async def test_summary_uses_model_without_browsing_tools(self) -> None:
        client = AsyncMock()
        client.responses.create.return_value.output_text = "Summary of research"
        context = AsyncMock()
        context.__aenter__.return_value = client
        with patch.dict(os.environ, {"OPENAI_API_KEY": "test"}), \
                patch.object(source_processing.openai, "AsyncOpenAI", return_value=context):
            self.assertEqual(await source_processing.summarize("Research facts"), "Summary of research")
        self.assertNotIn("tools", client.responses.create.call_args.kwargs)

    async def test_real_docling_manifest_with_injected_client_never_calls_network(self) -> None:
        from types import SimpleNamespace
        from docling_core.types.doc import DoclingDocument, DocItemLabel
        import docxtractor
        document = DoclingDocument(name="research")
        document.add_text(label=DocItemLabel.PARAGRAPH, text="Research results")
        client = Mock()
        client.convert.return_value = SimpleNamespace(status="success", document=document, errors=[])
        with tempfile.TemporaryDirectory() as directory:
            result = docxtractor.extract_document(
                b"# Research results", filename="paper.md", output_dir=directory,
                client=client,
            )
        self.assertEqual(result["source_texts"][0]["text"], "Research results")
        options = client.convert.call_args.kwargs["options"]
        self.assertTrue(options.do_picture_description)
        self.assertTrue(options.do_formula_enrichment)
        self.assertTrue(options.do_chart_extraction)
        self.assertEqual(options.image_export_mode, "referenced")
        self.assertTrue(options.include_images)
        self.assertTrue(options.include_page_images)

    def test_referenced_docling_bundle_preserves_extracted_pngs(self) -> None:
        from types import SimpleNamespace
        from PIL import Image
        from docling.service_client import DoclingServiceClient
        from docling_core.types.doc import DoclingDocument, ImageRef
        import docxtractor

        document = DoclingDocument(name="research")
        picture = document.add_picture(
            image=ImageRef.from_pil(Image.new("RGB", (320, 400), "red"), dpi=150),
        )
        picture.image.uri = Path("artifacts/picture.png")
        png = io.BytesIO()
        Image.new("RGB", (320, 400), "red").save(png, format="PNG")
        bundle = io.BytesIO()
        with zipfile.ZipFile(bundle, "w") as archive:
            archive.writestr("document.json", document.model_dump_json())
            archive.writestr("artifacts/picture.png", png.getvalue())

        with DoclingServiceClient(url="https://docling.example") as sdk:
            restored = sdk._document_from_artifact("resource_bundle", bundle.getvalue())
        client = Mock()
        client.convert.return_value = SimpleNamespace(status="success", document=restored, errors=[])
        with tempfile.TemporaryDirectory() as directory:
            result = docxtractor.extract_document(
                b"# Research", filename="paper.md", output_dir=directory, client=client,
            )
            asset = Path(result["output_directory"]) / result["elements"][0]["asset"]
            with Image.open(asset) as image:
                self.assertEqual(image.size, (320, 400))
                self.assertEqual(image.format, "PNG")
        self.assertEqual(result["warnings"], [])


@unittest.skipUnless(os.environ.get("POSTERS_TEST_DATABASE_URL"), "Set POSTERS_TEST_DATABASE_URL for PostgreSQL tests")
class ContentDatabaseTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        test_posters.PosterDatabaseTests.setUpClass.__func__(cls)
        cls.store = content.SourceStore("POSTERS_TEST_DATABASE_URL", cls.schema)

    @classmethod
    def cleanup_schema(cls) -> None:
        test_posters.PosterDatabaseTests.cleanup_schema.__func__(cls)

    def setUp(self) -> None:
        test_posters.PosterDatabaseTests.setUp(self)

    def test_multiple_top_level_sources_persist_in_order_and_are_owner_scoped(self) -> None:
        pid, _ = asyncio.run(self.store.new_or_existing(7))
        with tempfile.TemporaryDirectory() as directory, patch.dict(os.environ, {"SOURCES_DIR": directory, "GCS_BUCKET": ""}), \
                patch.object(content.processing, "extract", side_effect=lambda name, data: content.processing.Processed(text=data.decode())), \
                patch.object(content.processing, "summarize", AsyncMock(return_value="Summary")):
            first = asyncio.run(self.store.add(7, pid, "first.md", b"first"))
            second = asyncio.run(self.store.add(7, pid, "second.md", b"second"))
            self.assertEqual([item.sid for item in asyncio.run(self.store.list_sources(7, pid))], [first.sid, second.sid])
            self.assertEqual(asyncio.run(self.store.list_sources(8, pid)), [])
            rows = self.connection.execute("SELECT parent, storage_uri FROM source WHERE pid = %s", (pid,)).fetchall()
            self.assertTrue(all(parent is None and Path(uri).is_file() for parent, uri in rows))
            self.connection.execute((ROOT / "schema/schema.sql").read_text())


if __name__ == "__main__":
    unittest.main()
